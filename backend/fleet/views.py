import csv
import hashlib
import hmac
import io
import uuid
from datetime import datetime, timedelta, time
from decimal import Decimal
from django.conf import settings
from django.contrib.auth import authenticate
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth.tokens import default_token_generator
from django.core.mail import send_mail
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from django.db.models import OuterRef, Subquery, DecimalField
from django.db.models.functions import Coalesce
from django.http import HttpResponse, FileResponse
from django.middleware.csrf import get_token
from django.utils import timezone
from django.utils.http import urlsafe_base64_decode, urlsafe_base64_encode
from django.utils.encoding import force_bytes
from django.views.static import serve
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError, PermissionDenied
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework_simplejwt.tokens import AccessToken
from .authentication import enforce_csrf
from .models import User, Equipment, Reading, Maintenance, Supply, Alert, AuditLog, Sensor
from .permissions import RolePermission, MATRIX
from .serializers import UserSerializer, EquipmentSerializer, ReadingSerializer, MaintenanceSerializer, SupplySerializer, AlertSerializer, AuditSerializer, SensorSerializer
from .services import audit, reconcile_alert, LABELS
from .throttles import DeviceTelemetryThrottle

from django.utils.decorators import method_decorator
from django.views.decorators.cache import never_cache


@method_decorator(never_cache, name='dispatch')
class CSRFView(APIView):
    permission_classes = [AllowAny]
    authentication_classes = []
    def get(self, request): return Response({'csrfToken': get_token(request)})

class LoginView(APIView):
    permission_classes = [AllowAny]
    authentication_classes = []
    @transaction.atomic
    def post(self, request):
        enforce_csrf(request)
        email = str(request.data.get('email', '')).strip().lower()
        password = request.data.get('password', '')
        if not isinstance(password, str) or len(password) > 1024:
            return Response({'detail': 'Credenciales incorrectas.'}, status=401)
        user = User.objects.select_for_update().filter(email__iexact=email).first()
        if user and user.locked_until and user.locked_until > timezone.now():
            return Response({'detail': 'Acceso bloqueado temporalmente. Intenta nuevamente en 15 minutos.'}, status=429)
        if user and user.locked_until:
            user.failed_attempts = 0
            user.locked_until = None
            user.save(update_fields=['failed_attempts', 'locked_until'])
        authenticated = authenticate(request, username=user.username if user else email, password=password)
        if not authenticated:
            if user:
                user.failed_attempts += 1
                if user.failed_attempts >= 5: user.locked_until = timezone.now() + timedelta(minutes=15)
                user.save(update_fields=['failed_attempts', 'locked_until'])
            return Response({'detail': 'Correo o contraseña incorrectos.'}, status=401)
        user.failed_attempts, user.locked_until, user.last_login = 0, None, timezone.now()
        user.save(update_fields=['failed_attempts', 'locked_until', 'last_login'])
        response = Response(UserSerializer(user).data)
        response.set_cookie('access_token', str(AccessToken.for_user(user)), httponly=True, secure=not settings.DEBUG, samesite='Lax', max_age=7200)
        audit(user, 'login', user)
        return response

class LogoutView(APIView):
    permission_classes = [AllowAny]
    authentication_classes = []
    def post(self, request):
        enforce_csrf(request)
        response = Response({'detail': 'Sesión cerrada.'})
        response.delete_cookie('access_token', samesite='Lax')
        return response

class MeView(APIView):
    def get(self, request):
        data = UserSerializer(request.user).data
        data['permissions'] = {module: roles.get(request.user.role, '') for module, roles in MATRIX.items()}
        return Response(data)

class PasswordResetView(APIView):
    permission_classes = [AllowAny]
    authentication_classes = []
    def post(self, request):
        enforce_csrf(request)
        user = User.objects.filter(email__iexact=str(request.data.get('email', '')).strip(), is_active=True).first()
        if user:
            uid = urlsafe_base64_encode(force_bytes(user.pk))
            token = default_token_generator.make_token(user)
            url = f'{settings.FRONTEND_URL}/?reset={uid}&token={token}'
            send_mail('Recupera tu acceso · Vila’s Motor', f'Restablece tu contraseña en los próximos 30 minutos:\n{url}', settings.DEFAULT_FROM_EMAIL, [user.email])
        return Response({'detail': 'Si el correo está registrado, recibirás un enlace para recuperar tu acceso.'})

class PasswordResetConfirmView(APIView):
    permission_classes = [AllowAny]
    authentication_classes = []
    def post(self, request):
        enforce_csrf(request)
        try:
            user = User.objects.get(pk=urlsafe_base64_decode(request.data.get('uid', '')).decode())
        except (User.DoesNotExist, ValueError, UnicodeDecodeError):
            raise ValidationError('Enlace inválido o vencido.')
        if not default_token_generator.check_token(user, request.data.get('token', '')):
            raise ValidationError('Enlace inválido o vencido.')
        password = request.data.get('password', '')
        if not isinstance(password, str): raise ValidationError('Contraseña inválida.')
        try: validate_password(password, user)
        except DjangoValidationError as exc: raise ValidationError({'password': list(exc.messages)})
        user.set_password(password)
        user.failed_attempts, user.locked_until = 0, None
        user.save()
        audit(user, 'password_reset', user)
        return Response({'detail': 'Contraseña actualizada. Ya puedes ingresar.'})

class AuditedViewSet(viewsets.ModelViewSet):
    permission_classes = [RolePermission]
    def perform_create(self, serializer):
        instance = serializer.save()
        audit(self.request.user, 'created', instance)
    def perform_update(self, serializer):
        instance = serializer.save()
        audit(self.request.user, 'updated', instance)

class EquipmentViewSet(AuditedViewSet):
    queryset = Equipment.objects.annotate(
        _current_value=Coalesce(Subquery(Reading.objects.filter(equipment=OuterRef('pk')).order_by('-recorded_at', '-id').values('value')[:1]), 'initial_value', output_field=DecimalField()),
        _last_service_value=Coalesce(Subquery(Maintenance.objects.filter(equipment=OuterRef('pk'), status='completed').order_by('-value', '-id').values('value')[:1]), 'initial_service_value', output_field=DecimalField()))
    serializer_class = EquipmentSerializer
    permission_module = 'equipment'
    @transaction.atomic
    def perform_create(self, serializer):
        super().perform_create(serializer)
        reconcile_alert(serializer.instance, self.request.user)
    @transaction.atomic
    def update(self, request, *args, **kwargs):
        Equipment.objects.select_for_update().get(pk=kwargs['pk'])
        return super().update(request, *args, **kwargs)
    def perform_update(self, serializer):
        super().perform_update(serializer)
        instance = serializer.instance
        for key in ('_current_value', '_last_service_value', 'current_value', 'last_service_value', 'next_service', 'remaining', 'progress', 'status'):
            instance.__dict__.pop(key, None)
        reconcile_alert(serializer.instance, self.request.user)
    @transaction.atomic
    def perform_destroy(self, instance):
        Equipment.objects.select_for_update().get(pk=instance.pk)
        instance.active = False
        instance.save(update_fields=['active'])
        audit(self.request.user, 'deactivated', instance)
        reconcile_alert(instance, self.request.user)
    @action(detail=True, methods=['get'])
    def history(self, request, pk=None):
        equipment = self.get_object()
        return Response({'equipment': self.get_serializer(equipment).data,
            'readings': ReadingSerializer(equipment.readings.all(), many=True, context={'request': request}).data,
            'maintenances': MaintenanceSerializer(equipment.maintenances.all(), many=True).data,
            'alerts': AlertSerializer(equipment.alerts.all(), many=True).data,
            'sensors': SensorSerializer(equipment.sensors.all(), many=True).data})

class ReadingViewSet(AuditedViewSet):
    queryset = Reading.objects.select_related('equipment', 'created_by').all()
    serializer_class = ReadingSerializer
    permission_module = 'reading'
    def get_queryset(self):
        qs = super().get_queryset()
        if self.request.query_params.get('equipment'): qs = qs.filter(equipment_id=self.request.query_params['equipment'])
        return qs
    @transaction.atomic
    def create(self, request, *args, **kwargs):
        equipment_id = request.data.get('equipment')
        try: Equipment.objects.select_for_update().get(pk=equipment_id)
        except (Equipment.DoesNotExist, ValueError, TypeError): raise ValidationError({'equipment': 'Equipo inválido.'})
        client_id = request.data.get('client_id')
        if client_id:
            try: client_id = uuid.UUID(str(client_id))
            except ValueError: raise ValidationError({'client_id': 'Identificador inválido.'})
            existing = Reading.objects.filter(client_id=client_id).first()
            if existing:
                if existing.created_by_id != request.user.pk or str(existing.equipment_id) != str(equipment_id):
                    raise PermissionDenied('El identificador pertenece a otro registro.')
                return Response(self.get_serializer(existing).data, status=200)
        return super().create(request, *args, **kwargs)
    def perform_create(self, serializer):
        reading = serializer.save(created_by=self.request.user, source='manual')
        audit(self.request.user, 'created', reading)
        reconcile_alert(reading.equipment, self.request.user)
    @transaction.atomic
    def update(self, request, *args, **kwargs):
        reading = self.get_object()
        Equipment.objects.select_for_update().get(pk=reading.equipment_id)
        return super().update(request, *args, **kwargs)
    def perform_update(self, serializer):
        super().perform_update(serializer)
        reconcile_alert(serializer.instance.equipment, self.request.user)
    @transaction.atomic
    def perform_destroy(self, instance):
        Equipment.objects.select_for_update().get(pk=instance.equipment_id)
        equipment = instance.equipment
        if equipment.maintenances.filter(status='completed', value__gt=equipment.initial_value).exists():
            raise ValidationError('Esta lectura forma parte de un historial con mantenciones. Conserva el registro y corrige su valor.')
        audit(self.request.user, 'deleted', instance)
        instance.delete()
        reconcile_alert(equipment, self.request.user)

class MaintenanceViewSet(AuditedViewSet):
    queryset = Maintenance.objects.select_related('equipment').prefetch_related('supplies__supply').all()
    serializer_class = MaintenanceSerializer
    permission_module = 'maintenance'
    @transaction.atomic
    def create(self, request, *args, **kwargs):
        try: Equipment.objects.select_for_update().get(pk=request.data.get('equipment'))
        except (Equipment.DoesNotExist, ValueError, TypeError): raise ValidationError({'equipment': 'Equipo inválido.'})
        return super().create(request, *args, **kwargs)
    def perform_create(self, serializer):
        instance = serializer.save(created_by=self.request.user)
        audit(self.request.user, 'created', instance)
        reconcile_alert(instance.equipment, self.request.user, instance if instance.status == 'completed' else None)
    @transaction.atomic
    def update(self, request, *args, **kwargs):
        Equipment.objects.select_for_update().get(pk=self.get_object().equipment_id)
        return super().update(request, *args, **kwargs)
    def perform_update(self, serializer):
        super().perform_update(serializer)
        reconcile_alert(serializer.instance.equipment, self.request.user, serializer.instance)
    @transaction.atomic
    def perform_destroy(self, instance):
        Equipment.objects.select_for_update().get(pk=instance.equipment_id)
        instance.status = 'cancelled'
        instance.save(update_fields=['status'])
        audit(self.request.user, 'cancelled', instance)
        reconcile_alert(instance.equipment, self.request.user)

class UserViewSet(AuditedViewSet):
    queryset = User.objects.all().order_by('first_name')
    serializer_class = UserSerializer
    permission_module = 'user'
    def perform_destroy(self, instance):
        if instance.pk == self.request.user.pk: raise ValidationError('No puedes desactivar tu propia cuenta.')
        instance.is_active = False
        instance.save(update_fields=['is_active'])
        audit(self.request.user, 'deactivated', instance)

class SupplyViewSet(AuditedViewSet):
    queryset = Supply.objects.all().order_by('code')
    serializer_class = SupplySerializer
    permission_module = 'supply'
    def perform_destroy(self, instance):
        instance.active = False
        instance.save(update_fields=['active'])
        audit(self.request.user, 'deactivated', instance)

class AlertViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = Alert.objects.select_related('equipment').all()
    serializer_class = AlertSerializer
    permission_classes = [RolePermission]
    permission_module = 'alert'

class AuditViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = AuditLog.objects.select_related('user').all()[:1000]
    serializer_class = AuditSerializer
    permission_classes = [RolePermission]
    permission_module = 'audit'

class TelemetryView(APIView):
    permission_classes = [AllowAny]
    authentication_classes = []
    throttle_classes = [DeviceTelemetryThrottle]
    @transaction.atomic
    def post(self, request):
        token = request.headers.get('Authorization', '')
        if not token.startswith('Device '): return Response({'detail': 'Token de dispositivo requerido.'}, status=401)
        token_hash = hashlib.sha256(token[7:].encode()).hexdigest()
        sensor = Sensor.objects.select_related('equipment').filter(token_hash=token_hash, active=True).first()
        if not sensor: return Response({'detail': 'Dispositivo no autorizado.'}, status=401)
        Equipment.objects.select_for_update().get(pk=sensor.equipment_id)
        data = request.data.copy()
        data['equipment'] = sensor.equipment_id
        client_id = data.get('client_id')
        if client_id:
            try: uuid.UUID(str(client_id))
            except ValueError: raise ValidationError({'client_id': 'Identificador inválido.'})
            existing = Reading.objects.filter(client_id=client_id).first()
            if existing:
                if existing.sensor_id != sensor.pk: raise PermissionDenied('Identificador de lectura ya utilizado.')
                return Response(ReadingSerializer(existing, context={'request': request}).data)
        serializer = ReadingSerializer(data=data, context={'request': request})
        serializer.is_valid(raise_exception=True)
        reading = serializer.save(sensor=sensor, created_by=None, source='ocr' if data.get('confidence') is not None else 'sensor')
        if sensor.gateway:
            sensor.gateway.last_connection = timezone.now()
            sensor.gateway.save(update_fields=['last_connection'])
        audit(None, 'telemetry_received', reading)
        reconcile_alert(sensor.equipment)
        return Response(serializer.data, status=201)

def report_rows(day):
    cutoff = timezone.make_aware(datetime.combine(day, time.max))
    rows = []
    equipment = Equipment.objects.filter(active=True).annotate(
        report_value=Coalesce(Subquery(Reading.objects.filter(equipment=OuterRef('pk'), recorded_at__lte=cutoff).order_by('-recorded_at', '-id').values('value')[:1]), 'initial_value', output_field=DecimalField()),
        report_baseline=Coalesce(Subquery(Maintenance.objects.filter(equipment=OuterRef('pk'), status='completed', completed_at__lte=cutoff).order_by('-value', '-id').values('value')[:1]), 'initial_service_value', output_field=DecimalField()))
    for eq in equipment:
        value = eq.report_value
        baseline = eq.report_baseline
        next_value = baseline + eq.interval
        remaining = next_value - value
        state = 'overdue' if remaining < 0 else 'required' if remaining == 0 else 'upcoming' if remaining <= Decimal(eq.interval)*Decimal('0.15') else 'operational'
        rows.append({'id': eq.id, 'code': eq.code, 'name': eq.name, 'current_value': str(value), 'next_service': str(next_value),
                     'remaining': str(remaining), 'unit': eq.unit, 'status': state, 'status_label': LABELS[state]})
    return rows

class ReportView(APIView):
    def get(self, request):
        try: day = datetime.strptime(request.query_params.get('date', timezone.localdate().isoformat()), '%Y-%m-%d').date()
        except ValueError: raise ValidationError('Fecha inválida. Usa AAAA-MM-DD.')
        rows = report_rows(day)
        output = request.query_params.get('format', 'json')
        # DRF reserves ?format for renderer negotiation: exports use ?export.
        output = request.query_params.get('export', output)
        headers = ['Equipo', 'Nombre', 'Lectura actual', 'Próxima mantención', 'Restante', 'Unidad', 'Estado']
        values = [[r['code'], r['name'], float(r['current_value']), float(r['next_service']), float(r['remaining']), r['unit'], r['status_label']] for r in rows]
        if output == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
            wb = Workbook(); ws = wb.active; ws.title = 'Estado de la flota'
            ws.append([f"Vila's Motor · Informe diario · {day}"]); ws.append(headers)
            for row in values:
                ws.append([ILLEGAL_CHARACTERS_RE.sub('', v) if isinstance(v, str) else v for v in row])
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str): cell.data_type = 's'
            for cell in ws[2]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='173F2B')
            for col in 'ABCDEFG': ws.column_dimensions[col].width = 26
            buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
            return FileResponse(buffer, as_attachment=True, filename=f'flota-{day}.xlsx')
        if output == 'pdf':
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from xml.sax.saxutils import escape
            buffer = io.BytesIO(); styles = getSampleStyleSheet()
            doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=25, rightMargin=25)
            table_values = [headers] + [[Paragraph(escape(str(v)), styles['BodyText']) for v in row] for row in values]
            table = Table(table_values, colWidths=[60, 145, 105, 115, 75, 45, 185], repeatRows=1)
            table.setStyle(TableStyle([('BACKGROUND', (0,0),(-1,0), colors.HexColor('#173F2B')), ('TEXTCOLOR',(0,0),(-1,0),colors.white),
                ('VALIGN',(0,0),(-1,-1),'TOP'), ('FONTSIZE',(0,0),(-1,-1),8), ('BOTTOMPADDING',(0,0),(-1,-1),9),
                ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white, colors.HexColor('#F5F7F8')]), ('GRID',(0,0),(-1,-1),0.3,colors.lightgrey)]))
            doc.build([Paragraph(f"Vila's Motor · Informe diario de máquinas · {day}", styles['Heading1']), Spacer(1,16), table])
            buffer.seek(0)
            return FileResponse(buffer, as_attachment=True, filename=f'flota-{day}.pdf')
        return Response({'date': day.isoformat(), 'equipment': rows})

class MigrationTemplateView(APIView):
    def get(self, request):
        buffer = io.StringIO()
        csv.writer(buffer).writerow(['code', 'name', 'equipment_type', 'brand', 'model', 'serial_number', 'region', 'commune', 'location', 'responsible', 'unit', 'interval', 'initial_value', 'initial_service_value'])
        response = HttpResponse('\ufeff'+buffer.getvalue(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="plantilla-equipos.csv"'
        return response

class ProtectedMediaView(APIView):
    def get(self, request, path): return serve(request, path, document_root=settings.MEDIA_ROOT)
